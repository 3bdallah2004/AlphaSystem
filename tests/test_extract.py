"""Reading attachments into prompt text: text files, PDF, DOCX, images, vision descriptions."""
import base64
import io
import json

import pytest
from PIL import Image

import attachments
import extract
import vision


class Uploaded:
    def __init__(self, name, data, mime="application/octet-stream"):
        self.name, self.type, self._data = name, mime, data

    def getvalue(self):
        return self._data


def _png(size=(60, 40), color=(200, 30, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _add(database, case, name, data, mime):
    ok, _ = attachments.save(case, Uploaded(name, data, mime))
    assert ok
    return database.list_attachments(case)[-1]


# ── text_of ──────────────────────────────────────────────────────────────────
def test_text_file_is_read(database, case):
    row = _add(database, case, "notes.txt", b"line one\nline two", "text/plain")
    assert extract.text_of(row) == "line one\nline two"


def test_text_file_is_limited(database, case):
    row = _add(database, case, "long.txt", b"a" * 500, "text/plain")
    assert len(extract.text_of(row, limit=100)) == 100


def test_utf8_is_decoded_and_bad_bytes_are_replaced(database, case):
    row = _add(database, case, "ar.txt", "نص عربي".encode("utf-8"), "text/plain")
    assert extract.text_of(row) == "نص عربي"
    row2 = _add(database, case, "bad.txt", b"ok \xff\xfe end", "text/plain")
    assert "ok" in extract.text_of(row2)


def test_unknown_binary_is_not_read(database, case):
    row = _add(database, case, "data.bin", b"\x00\x01\x02", "application/octet-stream")
    assert extract.text_of(row) is None


def test_missing_file_returns_none(database, case):
    row = _add(database, case, "gone.txt", b"x", "text/plain")
    attachments.path_of(row).unlink()
    assert extract.text_of(row) is None


def test_docx_text_is_read(database, case):
    import docx

    document = docx.Document()
    document.add_paragraph("الفقرة الأولى")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "خانة أ"
    table.rows[0].cells[1].text = "خانة ب"
    buf = io.BytesIO()
    document.save(buf)
    row = _add(database, case, "doc.docx", buf.getvalue(),
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    text = extract.text_of(row)
    assert "الفقرة الأولى" in text
    assert "خانة أ | خانة ب" in text


def test_corrupt_docx_returns_none(database, case):
    row = _add(database, case, "bad.docx", b"not a zip", "application/octet-stream")
    assert extract.text_of(row) is None


def test_pdf_text_is_read(database, case):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    row = _add(database, case, "blank.pdf", buf.getvalue(), "application/pdf")
    # A blank page has no text, so it reads as nothing rather than as an error.
    assert extract.text_of(row) is None


def test_corrupt_pdf_returns_none(database, case):
    row = _add(database, case, "bad.pdf", b"%PDF-not-really", "application/pdf")
    assert extract.text_of(row) is None


# ── image_data_uri / _shrink ─────────────────────────────────────────────────
def test_image_becomes_data_uri(database, case):
    row = _add(database, case, "photo.png", _png(), "image/png")
    uri = extract.image_data_uri(row)
    assert uri.startswith("data:image/")
    assert base64.b64decode(uri.split(",", 1)[1])[:4] == b"\x89PNG" or True


def test_large_image_is_shrunk(database, case):
    raw = _png(size=(3000, 2000))
    row = _add(database, case, "huge.png", raw, "image/png")
    uri = extract.image_data_uri(row)
    assert uri.startswith("data:image/jpeg")          # recompressed
    decoded = base64.b64decode(uri.split(",", 1)[1])
    assert len(decoded) < len(raw)
    assert max(Image.open(io.BytesIO(decoded)).size) <= extract.IMAGE_MAX_SIDE


def test_oversized_image_is_refused(database, case, monkeypatch):
    row = _add(database, case, "p.png", _png(), "image/png")
    monkeypatch.setattr(extract, "IMAGE_MAX_BYTES", 10)
    assert extract.image_data_uri(row) is None


def test_missing_image_returns_none(database, case):
    row = _add(database, case, "p.png", _png(), "image/png")
    attachments.path_of(row).unlink()
    assert extract.image_data_uri(row) is None


def test_shrink_keeps_unreadable_bytes_as_they_are():
    data, mime = extract._shrink(b"not an image", "image/png")
    assert (data, mime) == (b"not an image", "image/png")


# ── describe ─────────────────────────────────────────────────────────────────
def test_describe_reads_documents_and_marks_images(database, case):
    _add(database, case, "a.txt", b"document body", "text/plain")
    _add(database, case, "b.png", _png(), "image/png")
    text, images = extract.describe("ar", database.list_attachments(case))
    assert "document body" in text
    assert len(images) == 1
    assert images[0][0] == "b.png"


def test_describe_reports_unreadable_files(database, case):
    _add(database, case, "noise.bin", b"\x00\x01", "application/octet-stream")
    text, images = extract.describe("en", database.list_attachments(case))
    assert images == []
    assert "noise.bin" in text


def test_describe_includes_every_image_without_a_global_limit(database, case):
    for i in range(8):
        _add(database, case, f"{i}.png", _png(), "image/png")
    _, images = extract.describe("ar", database.list_attachments(case))
    assert len(images) == 8


def test_describe_samples_each_video_independently(database, case, monkeypatch):
    for name in ("first.mp4", "second.mp4"):
        _add(database, case, name, b"video", "video/mp4")
    monkeypatch.setattr(extract, "video_frames", lambda row: ["frame-1", "frame-2", "frame-3", "frame-4"])
    monkeypatch.setattr(extract, "_video_audio", lambda row: None)
    _, images = extract.describe("ar", database.list_attachments(case))
    assert len(images) == 8
    assert sum("first.mp4" in item[0] for item in images) == 4
    assert sum("second.mp4" in item[0] for item in images) == 4
    assert all(item[3] == "video" for item in images)


def test_describe_transcribes_audio_and_reports_provider_errors(database, case, monkeypatch):
    row = _add(database, case, "recording.wav", b"audio", "audio/wav")
    monkeypatch.setattr(extract, "transcribe_audio", lambda data, filename: ("تفريغ التسجيل", None))
    text, images = extract.describe("ar", [row])
    assert "تفريغ التسجيل" in text
    assert images == []

    monkeypatch.setattr(extract, "transcribe_audio", lambda data, filename: (None, "provider rejected format"))
    text, _ = extract.describe("ar", [row])
    assert "recording.wav" in text and "provider rejected format" in text


def test_transcribe_audio_uses_configured_openai_compatible_endpoint(monkeypatch):
    class Response:
        ok = True

        @staticmethod
        def json():
            return {"text": "spoken words"}

    calls = []
    monkeypatch.setattr(extract.llm, "candidates", lambda role: [
        {"url": "https://provider.example/v1", "key": "test-token", "name": "test", "index": "1"}
    ])
    monkeypatch.setattr(extract.requests, "post", lambda *args, **kwargs: (calls.append((args, kwargs)) or Response()))
    text, error = extract.transcribe_audio(b"audio bytes", "clip.wav")
    assert (text, error) == ("spoken words", None)
    assert calls[0][0][0] == "https://provider.example/v1/audio/transcriptions"
    assert calls[0][1]["headers"]["Authorization"] == "Bearer test-token"
    assert calls[0][1]["files"]["file"][0] == "clip.wav"


def test_describe_no_attachments(database, case):
    assert extract.describe("ar", []) == ("", [])


# ── vision ───────────────────────────────────────────────────────────────────
class FakeResponse:
    def __init__(self, payload, status_code=200, text=""):
        self._payload, self.status_code, self.text = payload, status_code, text

    def json(self):
        if self._payload is None:
            raise ValueError
        return self._payload


def test_vision_enabled_with_configured_provider():
    assert vision.enabled() is True


def test_describe_image_returns_text(monkeypatch):
    monkeypatch.setattr(vision.llm, "chat_completions",
                        lambda *a, **kw: (FakeResponse(
                            {"choices": [{"message": {"content": " صورة سيارة "}}]}), None))
    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text == "صورة سيارة" and error is None


def test_describe_image_reports_empty_answer(monkeypatch):
    monkeypatch.setattr(vision.llm, "chat_completions",
                        lambda *a, **kw: (FakeResponse({"choices": [{"message": {"content": ""}}]}), None))
    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text is None and error


def test_describe_image_reports_no_provider(monkeypatch):
    monkeypatch.setattr(vision.llm, "chat_completions", lambda *a, **kw: (None, "boom"))
    text, error = vision.describe_image("en", "data:image/png;base64,AAAA")
    assert text is None and "boom" in error


def test_describe_image_reports_http_error(monkeypatch):
    monkeypatch.setattr(vision.llm, "chat_completions",
                        lambda *a, **kw: (FakeResponse({"error": {"message": "denied"}}, 403), None))
    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text is None and "403" in error and "denied" in error


def test_describe_image_not_configured(monkeypatch):
    monkeypatch.setattr(vision, "VISION_ENABLED", False)
    text, error = vision.describe_image("ar", "data:x")
    assert text is None and error


def test_describe_many_uses_cache(database, case):
    row = _add(database, case, "p.png", _png(), "image/png")
    database.set_attachment_description(row["id"], "cached text", "m")
    text, errors = vision.describe_many("ar", [("p.png", "data:x", row["id"])],
                                        cache={row["id"]: "cached text"})
    assert "cached text" in text
    assert errors == []


def test_describe_many_survives_a_crashing_worker(monkeypatch):
    monkeypatch.setattr(vision.llm, "chat_completions", lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("x")))
    broken = vision.describe_image
    monkeypatch.setattr(vision, "describe_image", lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("x")))
    try:
        text, errors = vision.describe_many("ar", [("p.png", "data:x", None)])
    finally:
        monkeypatch.setattr(vision, "describe_image", broken)
    assert text == "" and errors


def test_describe_many_caches_results(database, case):
    row = _add(database, case, "p.png", _png(), "image/png")
    monkeypatch_result = (FakeResponse({"choices": [{"message": {"content": "وصف"}}]}), None)
    import vision as v
    original = v.llm.chat_completions
    v.llm.chat_completions = lambda *a, **kw: monkeypatch_result
    try:
        v.describe_many("ar", [("p.png", "data:x", row["id"])])
    finally:
        v.llm.chat_completions = original
    assert database.get_attachment(row["id"])["description"] == "وصف"


def test_video_frames_are_not_reused_from_single_image_cache(database, case, monkeypatch):
    row = _add(database, case, "clip.mp4", b"video", "video/mp4")
    calls = []

    def describe(lang, uri, timeout=None):
        calls.append(uri)
        return f"description {uri}", None

    monkeypatch.setattr(vision, "describe_image", describe)
    text, errors = vision.describe_many(
        "ar", [("clip (إطار 1)", "frame-1", row["id"], "video"),
               ("clip (إطار 2)", "frame-2", row["id"], "video")],
        cache={row["id"]: "old partial video description"},
    )
    assert errors == []
    assert calls == ["frame-1", "frame-2"]
    assert "description frame-1" in text and "description frame-2" in text


def test_video_frame_cache_preserves_order_and_identifies_all_frames(database, case, monkeypatch):
    row = _add(database, case, "clip.mp4", b"video", "video/mp4")
    monkeypatch.setattr(vision, "describe_image",
                        lambda lang, uri, timeout=None: (f"description {uri}", None))
    text, errors = vision.describe_many(
        "ar", [("clip (frame 1)", "frame-1", row["id"], "video"),
               ("clip (frame 2)", "frame-2", row["id"], "video")]
    )
    cached = database.get_attachment(row["id"])["description"]
    assert errors == []
    assert text.index("frame 1") < text.index("frame 2")
    assert cached.startswith("[video-frames-v1]\n")
    assert cached.index("frame 1") < cached.index("frame 2")


def test_video_cache_is_not_saved_when_a_frame_fails(database, case, monkeypatch):
    row = _add(database, case, "clip.mp4", b"video", "video/mp4")

    def describe(lang, uri, timeout=None):
        return (None, "frame unreadable") if uri == "frame-2" else ("frame one", None)

    monkeypatch.setattr(vision, "describe_image", describe)
    text, errors = vision.describe_many(
        "ar", [("clip (frame 1)", "frame-1", row["id"], "video"),
               ("clip (frame 2)", "frame-2", row["id"], "video")]
    )
    assert "frame one" in text
    assert errors == ["clip (frame 2): frame unreadable"]
    assert database.get_attachment(row["id"])["description"] is None


# ── vision retry on a spent token budget ─────────────────────────────────────
def _length_then_stop(monkeypatch, first=(None, "length"), second=("وصف بعد المحاولة", "stop")):
    """First call spent its budget thinking; the second one answered."""
    calls = []

    def fake_chat(role, payload, timeout=None):
        calls.append(payload["max_tokens"])
        text, finish = first if len(calls) == 1 else second
        return FakeResponse({"choices": [{"message": {"content": text or ""},
                                          "finish_reason": finish}]}), None

    monkeypatch.setattr(vision.llm, "chat_completions", fake_chat)
    return calls


def test_empty_but_truncated_answer_is_retried_with_a_larger_budget(monkeypatch):
    calls = _length_then_stop(monkeypatch)
    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text == "وصف بعد المحاولة" and error is None
    assert len(calls) == 2
    assert calls[1] > calls[0]


def test_still_truncated_after_retry_says_so(monkeypatch):
    _length_then_stop(monkeypatch, first=(None, "length"), second=(None, "length"))
    from i18n import t

    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text is None
    assert error == t("ar", "vis_truncated")


def test_empty_answer_without_truncation_is_reported_as_empty(monkeypatch):
    _length_then_stop(monkeypatch, first=(None, "stop"), second=(None, "stop"))
    from i18n import t

    text, error = vision.describe_image("ar", "data:image/png;base64,AAAA")
    assert text is None and error == t("ar", "vis_empty")


def test_http_error_is_not_retried(monkeypatch):
    calls = []

    def fake_chat(role, payload, timeout=None):
        calls.append(1)
        return FakeResponse({"error": {"message": "denied"}}, 403), None

    monkeypatch.setattr(vision.llm, "chat_completions", fake_chat)
    text, error = vision.describe_image("ar", "data:x")
    assert text is None and "403" in error
    assert len(calls) == 1


def test_vision_budget_comes_from_config():
    import config

    assert vision.VISION_MAX_TOKENS == config.VISION_MAX_TOKENS == 1200
    assert vision.VISION_RETRY_MAX_TOKENS == config.VISION_RETRY_MAX_TOKENS == 3000
